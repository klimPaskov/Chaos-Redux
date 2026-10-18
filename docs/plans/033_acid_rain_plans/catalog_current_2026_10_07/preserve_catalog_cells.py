import argparse
import copy
import hashlib
import html
import json
import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"

OLD_VALUES = {
    ("Events", "C34"): "A lethal atmospheric system moves through a frozen registry of valid states, touching each state exactly once while prepared countries spend resources to limit deaths, damage, contamination, and supply loss before the ordinary cloud dissipates.",
    ("Events", "D34"): "Severe Storm Cells: localized high-intensity cells add stronger human, infrastructure, supply, and military pressure while the regional system continues moving.",
    ("Events", "E34"): "Multiple Weather Fronts: two independent fronts can operate at once, with a conditional third front, separate warnings, timers, intensity, and severe cells.",
    ("Events", "F34"): "Global Acid Rain: simultaneous worldwide exposure closes the regional phase, applies the same exact population-loss contract everywhere, and retains a guaranteed dissipation path.",
    ("Events", "M34"): "To Be Reworked",
    ("Clusters", "C6"): "Natural and environmental catastrophes strike states or regions through five distinct Natural Disasters slots, followed by optional Acid Rain and Heat Wave crises. A sequence containing Acid Rain commits Major-event pacing once while preserving the cluster history.",
}

DESIRED_VALUES = {
    ("Events", "C34"): "A lethal acid-rain system moves between regions, exposing the inhabited states in its path to corrosion, poisoned runoff, infrastructure damage, and direct civilian loss. Countries can build shelters, protect water and food, expand emergency medicine, harden transport, evacuate states ahead of storm forecasts, and repair the aftermath. The rain ordinarily disperses only after it has crossed every state in its path.",
    ("Events", "D34"): "Severe Storm Cells: Concentrated cells form inside the moving front, giving threatened states nine days to prepare for heavier corrosion and lethal exposure.",
    ("Events", "E34"): "Multiple Weather Fronts: The atmospheric mass divides into independent fronts with separate routes, warnings, intensity, and severe cells.",
    ("Events", "F34"): "Global Acid Rain: The fronts merge into a worldwide layer that exposes every inhabited state in its path and produces regional superstorms with nine-day forecasts.",
    ("Events", "M34"): "Needs Testing",
    ("Clusters", "C6"): "Natural and environmental catastrophes strike states or regions, with repeated disaster slots allowing the cluster to scale into larger cascades.",
}


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def sheet_paths(archive):
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relations = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{NS_PACKAGE_REL}}}Relationship")
    }
    result = {}
    for sheet in workbook_root.findall(f".//{{{NS_MAIN}}}sheet"):
        relation_id = sheet.attrib[f"{{{NS_DOC_REL}}}id"]
        target = relations[relation_id]
        if target.startswith("/"):
            package_path = target.lstrip("/")
        else:
            package_path = posixpath.normpath(posixpath.join("xl", target))
        result[sheet.attrib["name"]] = package_path
    return result


def read_workbook_values(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        result = {}
        for key in DESIRED_VALUES:
            sheet_name, address = key
            result[key] = workbook[sheet_name][address].value
        return result
    finally:
        workbook.close()


def target_cell_pattern(address):
    encoded = re.escape(address.encode("ascii"))
    return re.compile(
        rb'<c\b(?=[^>]*\br="' + encoded + rb'")[^>]*>.*?</c>',
        re.DOTALL,
    )


def patch_cell_text(xml_bytes, address, value):
    pattern = target_cell_pattern(address)
    matches = list(pattern.finditer(xml_bytes))
    if len(matches) != 1:
        raise ValueError(f"Expected one XML cell {address}, found {len(matches)}")
    old_cell = matches[0].group(0)
    if b't="inlineStr"' not in old_cell or b"<is>" not in old_cell:
        raise ValueError(f"{address} is not a simple inline string cell")
    if b"<f" in old_cell:
        raise ValueError(f"{address} unexpectedly contains a formula")
    text_pattern = re.compile(rb"(<t\b[^>]*>)(.*?)(</t>)", re.DOTALL)
    text_matches = list(text_pattern.finditer(old_cell))
    if len(text_matches) != 1:
        raise ValueError(f"Expected one text payload in {address}, found {len(text_matches)}")
    escaped = html.escape(value, quote=False).encode("utf-8")
    new_cell = text_pattern.sub(lambda match: match.group(1) + escaped + match.group(3), old_cell, count=1)
    new_xml = xml_bytes[:matches[0].start()] + new_cell + xml_bytes[matches[0].end():]
    return new_xml, old_cell, new_cell


def zipinfo_signature(info):
    return (
        info.filename,
        info.date_time,
        info.compress_type,
        info.create_system,
        info.create_version,
        info.extract_version,
        info.flag_bits,
        info.volume,
        info.internal_attr,
        info.external_attr,
        info.comment,
        info.extra,
    )


def verify_cell_and_style_preservation(original_path, result_path):
    source = load_workbook(original_path, data_only=False)
    result = load_workbook(result_path, data_only=False)
    try:
        if source.sheetnames != result.sheetnames:
            raise ValueError("Worksheet names or order changed")
        value_changes = []
        style_changes = []
        formula_changes = []
        for sheet_name in source.sheetnames:
            old_sheet = source[sheet_name]
            new_sheet = result[sheet_name]
            if (old_sheet.max_row, old_sheet.max_column) != (new_sheet.max_row, new_sheet.max_column):
                raise ValueError(f"Used range changed on {sheet_name}")
            for row in range(1, old_sheet.max_row + 1):
                for col in range(1, old_sheet.max_column + 1):
                    old_cell = old_sheet.cell(row, col)
                    new_cell = new_sheet.cell(row, col)
                    key = (sheet_name, old_cell.coordinate)
                    if old_cell.value != new_cell.value:
                        value_changes.append(key)
                    if old_cell.style_id != new_cell.style_id:
                        style_changes.append(key)
                    if old_cell.data_type == "f" or new_cell.data_type == "f":
                        if old_cell.value != new_cell.value or old_cell.data_type != new_cell.data_type:
                            formula_changes.append(key)
            if old_sheet.freeze_panes != new_sheet.freeze_panes:
                raise ValueError(f"Freeze panes changed on {sheet_name}")
            if old_sheet.auto_filter.ref != new_sheet.auto_filter.ref:
                raise ValueError(f"AutoFilter changed on {sheet_name}")
            if str(old_sheet.merged_cells) != str(new_sheet.merged_cells):
                raise ValueError(f"Merged cells changed on {sheet_name}")
            if [(table.name, table.ref) for table in old_sheet.tables.values()] != [
                (table.name, table.ref) for table in new_sheet.tables.values()
            ]:
                raise ValueError(f"Tables changed on {sheet_name}")
            if old_sheet.sheet_view.showGridLines != new_sheet.sheet_view.showGridLines:
                raise ValueError(f"Gridline visibility changed on {sheet_name}")
            old_rows = {
                key: (dim.height, dim.hidden, dim.style, dim.outlineLevel, dim.collapsed)
                for key, dim in old_sheet.row_dimensions.items()
            }
            new_rows = {
                key: (dim.height, dim.hidden, dim.style, dim.outlineLevel, dim.collapsed)
                for key, dim in new_sheet.row_dimensions.items()
            }
            if old_rows != new_rows:
                raise ValueError(f"Row dimensions changed on {sheet_name}")
            old_columns = {
                key: (dim.width, dim.hidden, dim.style, dim.min, dim.max, dim.outlineLevel, dim.collapsed)
                for key, dim in old_sheet.column_dimensions.items()
            }
            new_columns = {
                key: (dim.width, dim.hidden, dim.style, dim.min, dim.max, dim.outlineLevel, dim.collapsed)
                for key, dim in new_sheet.column_dimensions.items()
            }
            if old_columns != new_columns:
                raise ValueError(f"Column dimensions changed on {sheet_name}")

        expected_changes = sorted(DESIRED_VALUES)
        if sorted(value_changes) != expected_changes:
            raise ValueError(f"Unexpected cell-value delta set: {value_changes}")
        if style_changes:
            raise ValueError(f"Cell styles changed: {style_changes}")
        if formula_changes:
            raise ValueError(f"Formula changes detected: {formula_changes}")
        if read_workbook_values(result_path) != DESIRED_VALUES:
            raise ValueError("Final target values do not match artifact-authored values")
        return {
            "changedCells": [f"{sheet}!{cell}" for sheet, cell in value_changes],
            "styleChanges": style_changes,
            "formulaChanges": formula_changes,
            "sheetNames": source.sheetnames,
        }
    finally:
        source.close()
        result.close()


def preserve(original_path, artifact_path, output_path):
    original_path = Path(original_path)
    artifact_path = Path(artifact_path)
    output_path = Path(output_path)
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing candidate: {output_path}")

    artifact_values = read_workbook_values(artifact_path)
    if artifact_values != DESIRED_VALUES:
        raise ValueError("Artifact-tool draft does not contain the six requested values")
    original_values = read_workbook_values(original_path)
    if original_values != OLD_VALUES:
        raise ValueError("Original workbook target cells changed from the reviewed baseline")

    with zipfile.ZipFile(original_path, "r") as source_zip:
        original_names = source_zip.namelist()
        paths = sheet_paths(source_zip)
        patched_sheets = {}
        old_cells = {}
        new_cells = {}
        for (sheet_name, address), value in DESIRED_VALUES.items():
            sheet_path = paths[sheet_name]
            xml_bytes = patched_sheets.get(sheet_path, source_zip.read(sheet_path))
            xml_bytes, old_cell, new_cell = patch_cell_text(xml_bytes, address, value)
            patched_sheets[sheet_path] = xml_bytes
            old_cells[(sheet_path, address)] = old_cell
            new_cells[(sheet_path, address)] = new_cell
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_path, "w", allowZip64=True) as result_zip:
            result_zip.comment = source_zip.comment
            for info in source_zip.infolist():
                data = patched_sheets.get(info.filename)
                if data is None:
                    data = source_zip.read(info.filename)
                result_zip.writestr(copy.copy(info), data)

    with zipfile.ZipFile(original_path, "r") as source_zip, zipfile.ZipFile(output_path, "r") as result_zip:
        if source_zip.namelist() != result_zip.namelist():
            raise ValueError("ZIP entry list or order changed")
        changed_members = []
        unchanged_members = []
        metadata_changes = []
        for old_info, new_info in zip(source_zip.infolist(), result_zip.infolist()):
            old_data = source_zip.read(old_info.filename)
            new_data = result_zip.read(new_info.filename)
            if zipinfo_signature(old_info) != zipinfo_signature(new_info):
                metadata_changes.append(old_info.filename)
            if old_data != new_data:
                changed_members.append(old_info.filename)
            else:
                unchanged_members.append(old_info.filename)
        expected_members = sorted({paths["Events"], paths["Clusters"]})
        if sorted(changed_members) != expected_members:
            raise ValueError(f"Unexpected changed ZIP payloads: {changed_members}")
        if metadata_changes:
            raise ValueError(f"ZIP member metadata changed: {metadata_changes}")

        for sheet_path in expected_members:
            normalized = result_zip.read(sheet_path)
            for (target_path, address), old_cell in old_cells.items():
                if target_path != sheet_path:
                    continue
                new_cell = new_cells[(target_path, address)]
                if normalized.count(new_cell) != 1:
                    raise ValueError(f"Expected one transplanted cell fragment {address}")
                normalized = normalized.replace(new_cell, old_cell, 1)
            if normalized != source_zip.read(sheet_path):
                raise ValueError(f"Non-target XML bytes changed in {sheet_path}")

    preservation = verify_cell_and_style_preservation(original_path, output_path)
    with zipfile.ZipFile(original_path, "r") as source_zip, zipfile.ZipFile(output_path, "r") as result_zip:
        members = {
            name: {
                "beforeSha256": sha256(source_zip.read(name)),
                "afterSha256": sha256(result_zip.read(name)),
                "unchanged": source_zip.read(name) == result_zip.read(name),
            }
            for name in source_zip.namelist()
        }
    return {
        "original": str(original_path.resolve()),
        "originalSha256": sha256(original_path.read_bytes()),
        "artifactDraft": str(artifact_path.resolve()),
        "artifactDraftSha256": sha256(artifact_path.read_bytes()),
        "candidate": str(output_path.resolve()),
        "candidateSha256": sha256(output_path.read_bytes()),
        "candidateBytes": output_path.stat().st_size,
        "preservation": preservation,
        "changedZipMembers": changed_members,
        "unchangedZipMemberCount": len(unchanged_members),
        "zipEntryCount": len(members),
        "zipMembers": members,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--original", required=True)
    parser.add_argument("--artifact", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    result = preserve(args.original, args.artifact, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
