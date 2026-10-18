import copy
import csv
import hashlib
import io
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
from preserve_catalog_cells import OLD_VALUES, patch_cell_text, sheet_paths, zipinfo_signature

ROOT = Path(__file__).resolve().parents[4]
CATALOG = ROOT / "docs/spreadsheets/chaos_redux_events_catalog.xlsx"
CSV_SPECS = {
    "Events": ("chaos_redux_events_catalog.csv", 13),
    "Clusters": ("chaos_redux_clusters_catalog.csv", 8),
    "Scenarios": ("chaos_redux_scenarios_catalog.csv", 6),
}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def git(*args):
    return subprocess.check_output(["git", *args], cwd=ROOT).decode("utf-8").strip()


def state():
    files = [CATALOG] + [CATALOG.parent / name for name, _ in CSV_SPECS.values()]
    return {
        "headCommit": git("rev-parse", "HEAD"),
        "indexEntries": git("ls-files", "-s", "--", *[str(p.relative_to(ROOT)) for p in files]).splitlines(),
        "fileSha256": {str(p.relative_to(ROOT)): digest(p.read_bytes()) for p in files},
    }


def locate(workbook, sheet_name, id_header, identity, fields):
    sheet = workbook[sheet_name]
    headers = []
    for row in sheet:
        values = [cell.value for cell in row]
        if id_header in values and all(field in values for field in fields):
            headers.append(row[0].row)
    assert len(headers) == 1, (sheet_name, "header rows", headers)
    header_row = headers[0]
    mapping = {cell.value: cell.column for cell in sheet[header_row] if cell.value is not None}
    matches = [
        row for row in range(header_row + 1, sheet.max_row + 1)
        if str(sheet.cell(row, mapping[id_header]).value).strip() == str(identity)
    ]
    assert len(matches) == 1, (sheet_name, identity, matches)
    return sheet, header_row, matches[0], mapping


def localisation(path, keys):
    result = {}
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        match = re.match(r'^([^#\s][^:]*):\s*"((?:\\.|[^"])*)"\s*$', line)
        if match and match.group(1) in keys:
            result[match.group(1)] = json.loads('"' + match.group(2) + '"')
    return result


def cell_delta(before, after):
    assert before.sheetnames == after.sheetnames
    values, styles, formulas = [], [], []
    for name in before.sheetnames:
        left, right = before[name], after[name]
        assert (left.max_row, left.max_column) == (right.max_row, right.max_column)
        for row in range(1, left.max_row + 1):
            for column in range(1, left.max_column + 1):
                a, b = left.cell(row, column), right.cell(row, column)
                if a.value != b.value:
                    values.append({"sheet": name, "cell": a.coordinate, "before": a.value, "after": b.value})
                if a.style_id != b.style_id:
                    styles.append(name + "!" + a.coordinate)
                if (a.data_type == "f" or b.data_type == "f") and (a.value != b.value or a.data_type != b.data_type):
                    formulas.append(name + "!" + a.coordinate)
    return {"values": values, "styles": styles, "formulas": formulas}


def feature_summary(workbook):
    result = {}
    for sheet in workbook:
        result[sheet.title] = {
            "state": sheet.sheet_state,
            "dimension": sheet.calculate_dimension(),
            "freezePanes": sheet.freeze_panes,
            "autoFilter": sheet.auto_filter.ref,
            "tables": {table.name: table.ref for table in sheet.tables.values()},
            "mergedRanges": str(sheet.merged_cells),
            "validationCount": len(sheet.data_validations.dataValidation),
            "conditionalFormattingCount": len(sheet.conditional_formatting),
            "rowDimensionsCount": len(sheet.row_dimensions),
            "rowDimensionsSha256": digest(json.dumps({
                str(key): [dim.height, dim.hidden, dim.style, dim.outlineLevel, dim.collapsed]
                for key, dim in sheet.row_dimensions.items()
            }, sort_keys=True).encode()),
            "columnDimensionsCount": len(sheet.column_dimensions),
            "columnDimensionsSha256": digest(json.dumps({
                key: [dim.width, dim.hidden, dim.style, dim.min, dim.max, dim.outlineLevel, dim.collapsed]
                for key, dim in sheet.column_dimensions.items()
            }, sort_keys=True).encode()),
            "showGridLines": sheet.sheet_view.showGridLines,
        }
    return result


def check_exports(workbook, directory):
    checks = {}
    for sheet_name, (filename, width) in CSV_SPECS.items():
        sheet = workbook[sheet_name]
        rows = [[value.isoformat() if hasattr(value, "isoformat") else "" if value is None else str(value)
                 for value in row]
                for row in sheet.iter_rows(min_col=1, max_col=width, values_only=True)]
        while rows and not any(rows[-1]):
            rows.pop()
        path = directory / filename
        with path.open(encoding="utf-8-sig", newline="") as stream:
            actual = list(csv.reader(stream))
        checks[sheet_name] = {
            "path": str(path), "rows": len(actual), "columns": width,
            "header": actual[0], "allRowWidthsCorrect": all(len(row) == width for row in actual),
            "matchesWorkbook": actual == rows, "sha256": digest(path.read_bytes()),
        }
        assert checks[sheet_name]["matchesWorkbook"] and checks[sheet_name]["allRowWidthsCorrect"]
    return checks


def main():
    before_state = state()
    current_raw = CATALOG.read_bytes()
    assert zipfile.is_zipfile(io.BytesIO(current_raw))
    pointer = git("show", before_state["headCommit"] + ":docs/spreadsheets/chaos_redux_events_catalog.xlsx")
    oid = re.search(r"oid sha256:([0-9a-f]{64})", pointer).group(1)
    object_root = Path(git("rev-parse", "--git-path", "lfs/objects"))
    if not object_root.is_absolute():
        object_root = ROOT / object_root
    lfs_path = object_root / oid[:2] / oid[2:4] / oid
    head_raw = lfs_path.read_bytes()
    assert digest(head_raw) == oid and zipfile.is_zipfile(io.BytesIO(head_raw))
    load = lambda raw: openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=True, rich_text=True)
    current, head = load(current_raw), load(head_raw)
    event_fields = ["Details", "Evo I", "Evo II", "Evo III", "Status"]
    current_event = locate(current, "Events", "ID", 33, event_fields)
    current_cluster = locate(current, "Clusters", "Cluster ID", 5, ["Details", "Members (ID)", "Member Severities"])
    head_event = locate(head, "Events", "ID", 33, event_fields)
    head_cluster = locate(head, "Clusters", "Cluster ID", 5, ["Details"])
    event_keys = {"acid_rain.event_details.description"} | {
        "acid_rain.evolution.stage_" + str(stage) + "." + field
        for stage in (1, 2, 3) for field in ("title", "body")
    }
    loc = localisation(ROOT / "localisation/english/033_acid_rain_l_english.yml", event_keys)
    gui_loc = localisation(ROOT / "localisation/english/chaosx_gui_l_english.yml",
                           {"chaosx.events_log.window.cluster_details.description.natural_disasters"})
    suffix = "[This.GetAcidRainEventDetailsStatus][This.GetAcidRainEventDetailsAchievements]"
    premise = loc["acid_rain.event_details.description"]
    assert premise.endswith(suffix)
    expected_fields = {
        "Details": premise[:-len(suffix)],
        "Evo I": loc["acid_rain.evolution.stage_1.title"] + ": " + loc["acid_rain.evolution.stage_1.body"],
        "Evo II": loc["acid_rain.evolution.stage_2.title"] + ": " + loc["acid_rain.evolution.stage_2.body"],
        "Evo III": loc["acid_rain.evolution.stage_3.title"] + ": " + loc["acid_rain.evolution.stage_3.body"],
        "Status": "Needs Testing",
    }
    desired = {
        ("Events", head_event[0].cell(head_event[2], head_event[3][field]).coordinate): value
        for field, value in expected_fields.items()
    }
    desired[("Clusters", head_cluster[0].cell(head_cluster[2], head_cluster[3]["Details"]).coordinate)] = \
        gui_loc["chaosx.events_log.window.cluster_details.description.natural_disasters"]
    current_targets = {
        "Events." + field: {
            "cell": current_event[0].cell(current_event[2], current_event[3][field]).coordinate,
            "value": current_event[0].cell(current_event[2], current_event[3][field]).value,
            "literal": current_event[0].cell(current_event[2], current_event[3][field]).data_type == "s",
        } for field in event_fields
    }
    current_targets["Clusters.Details"] = {
        "cell": current_cluster[0].cell(current_cluster[2], current_cluster[3]["Details"]).coordinate,
        "value": current_cluster[0].cell(current_cluster[2], current_cluster[3]["Details"]).value,
        "literal": current_cluster[0].cell(current_cluster[2], current_cluster[3]["Details"]).data_type == "s",
    }
    assert all(current_targets["Events." + field]["value"] == value for field, value in expected_fields.items())
    assert current_targets["Clusters.Details"]["value"] == desired[("Clusters", head_cluster[0].cell(head_cluster[2], head_cluster[3]["Details"]).coordinate)]
    assert all(entry["literal"] for entry in current_targets.values())
    for key in desired:
        assert head[key[0]][key[1]].value == OLD_VALUES[key]
    assert state() == before_state, "Source or Git state changed before isolated candidate creation"
    outside = ROOT.parent / ("chaos_redux_catalog_event033_head_" + before_state["headCommit"][:8] + "_resume")
    assert not outside.exists(), "Existing isolated output requires inspection, not an overwrite"
    outside.mkdir()
    (outside / "recorded_head_original.xlsx").write_bytes(head_raw)
    (outside / "current_observed_original.xlsx").write_bytes(current_raw)
    candidate = outside / "chaos_redux_events_catalog.xlsx"
    fragments, modified_parts = {}, {}
    with zipfile.ZipFile(io.BytesIO(head_raw)) as source, zipfile.ZipFile(candidate, "w") as output:
        paths = sheet_paths(source)
        for (name, address), value in desired.items():
            part = paths[name]
            data = modified_parts.get(part, source.read(part))
            style = re.search(rb'<c\b(?=[^>]*\br="' + address.encode() + rb'")[^>]*>', data).group()
            assert b't="inlineStr"' in style and b's="' + str(head[name][address].style_id).encode() + b'"' in style
            updated, old_cell, new_cell = patch_cell_text(data, address, value)
            modified_parts[part] = updated
            fragments[(part, address)] = (old_cell, new_cell)
        output.comment = source.comment
        for info in source.infolist():
            output.writestr(copy.copy(info), modified_parts.get(info.filename, source.read(info.filename)))
    completed = load(candidate.read_bytes())
    delta = cell_delta(head, completed)
    assert {(d["sheet"], d["cell"]) for d in delta["values"]} == set(desired)
    assert not delta["styles"] and not delta["formulas"]
    assert feature_summary(head) == feature_summary(completed)
    packages = {}
    with zipfile.ZipFile(io.BytesIO(head_raw)) as source, zipfile.ZipFile(candidate) as result:
        assert source.namelist() == result.namelist()
        for name in source.namelist():
            a, b = source.read(name), result.read(name)
            assert zipinfo_signature(source.getinfo(name)) == zipinfo_signature(result.getinfo(name))
            if name in modified_parts:
                restored = b
                for (part, _), (old_cell, new_cell) in fragments.items():
                    if part == name:
                        assert restored.count(new_cell) == 1
                        restored = restored.replace(new_cell, old_cell, 1)
                assert restored == a
            else:
                assert b == a
            packages[name] = {"beforeSha256": digest(a), "afterSha256": digest(b), "unchanged": a == b}
    export_command = [
        sys.executable, str(ROOT / ".tools/export_event_catalog_csv.py"),
        "--workbook", str(candidate), "--output-dir", str(outside / "exports"),
    ]
    export_run = subprocess.run(export_command, cwd=ROOT, text=True, capture_output=True)
    assert export_run.returncode == 0, export_run.stderr
    memberships = current["Cluster Memberships"]
    membership_header_row = next(row[0].row for row in memberships if "Slot" in [c.value for c in row] and "Event ID" in [c.value for c in row])
    membership_headers = {c.value: c.column for c in memberships[membership_header_row] if c.value is not None}
    members = []
    clusters_for_event33 = set()
    for row in range(membership_header_row + 1, memberships.max_row + 1):
        cid = memberships.cell(row, membership_headers["Cluster ID"]).value
        eid = memberships.cell(row, membership_headers["Event ID"]).value
        if eid == 33:
            clusters_for_event33.add(str(cid))
        if cid == 5:
            members.append({
                "row": row, "slot": memberships.cell(row, membership_headers["Slot"]).value,
                "eventId": eid, "severity": memberships.cell(row, membership_headers["Severity"]).value,
                "notes": memberships.cell(row, membership_headers["Membership Notes"]).value,
            })
    members.sort(key=lambda row: row["slot"])
    aggregate_ids = [int(v.strip()) for v in current_cluster[0].cell(current_cluster[2], current_cluster[3]["Members (ID)"]).value.split(",")]
    aggregate_severity = [v.strip() for v in current_cluster[0].cell(current_cluster[2], current_cluster[3]["Member Severities"]).value.split(",")]
    assert aggregate_ids == [row["eventId"] for row in members] == [13, 13, 13, 13, 13, 33, 51]
    assert aggregate_severity == [row["severity"] for row in members]
    assert clusters_for_event33 == set(str(current_event[0].cell(current_event[2], current_event[3]["Cluster ID"]).value).split(","))
    legend = current["Legend"]
    legend_statuses = [{"row": row[0].row, "values": [c.value for c in row], "fills": [
        {"pattern": c.fill.patternType, "colorType": c.fill.fgColor.type, "rgb": str(c.fill.fgColor.rgb),
         "indexed": str(c.fill.fgColor.indexed), "theme": str(c.fill.fgColor.theme), "tint": c.fill.fgColor.tint}
        for c in row]}
                       for row in legend if any(c.value in ("Needs Testing", "To Be Reworked", "Playable") for c in row)]
    status_cell = current_event[0].cell(current_event[2], current_event[3]["Status"])
    status_validations = [{"sqref": str(rule.sqref), "formula1": rule.formula1}
                          for rule in current_event[0].data_validations.dataValidation if status_cell.coordinate in rule.sqref]
    formulas = [sheet.title + "!" + c.coordinate for sheet in current for row in sheet for c in row if c.data_type == "f"]
    errors = [sheet.title + "!" + c.coordinate for sheet in current for row in sheet for c in row if c.data_type == "e"]
    cached_current = openpyxl.load_workbook(io.BytesIO(current_raw), read_only=True, data_only=True)
    cached_candidate = openpyxl.load_workbook(candidate, read_only=True, data_only=True)
    report = {
        "workflow": "chaos-redux-xlsx; read-only openpyxl plus six-node native OOXML candidate",
        "python": sys.executable, "openpyxlVersion": openpyxl.__version__, "stateAtCapture": before_state,
        "headLfsPointer": pointer, "headRawSha256": digest(head_raw), "headRawBytes": len(head_raw),
        "currentTargetsLocatedByIdAndHeader": current_targets, "liveWorkbookAlreadySatisfied": True,
        "noLiveWorkbookSaveOrExportOnResume": True, "headVersusLiveCellDelta": cell_delta(head, current),
        "isolatedDirectory": str(outside), "candidate": str(candidate), "candidateSha256": digest(candidate.read_bytes()),
        "candidateBytes": candidate.stat().st_size, "candidateCellDelta": delta,
        "nativeFeatureComparisonEqual": True, "headZipPayloads": packages,
        "currentFeatures": feature_summary(current), "clusterMemberships": members,
        "legendStatusRows": legend_statuses, "statusValidations": status_validations,
        "currentFormulas": formulas, "actualErrorCells": errors,
        "currentExports": check_exports(cached_current, CATALOG.parent),
        "isolatedExporterCommand": export_command, "isolatedExporterResult": json.loads(export_run.stdout),
        "isolatedExports": check_exports(cached_candidate, outside / "exports"),
        "stateAfterVerification": state(),
    }
    report["sourceAndGitStateStable"] = report["stateAtCapture"] == report["stateAfterVerification"]
    for book in (current, head, completed, cached_current, cached_candidate):
        book.close()
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
