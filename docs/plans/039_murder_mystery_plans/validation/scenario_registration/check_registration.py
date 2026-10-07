"""Read-only source-contract checks; not HOI4 runtime acceptance."""
from pathlib import Path
import re
import csv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[5]
BASE = Path(__file__).parent


def read(path):
    return (ROOT / path).read_text(encoding="utf-8-sig")


def block(text, name):
    match = re.search(r"(?m)^" + re.escape(name) + r"\s*=\s*\{", text)
    assert match, name
    start = match.end()
    depth = 1
    for end in range(start, len(text)):
        depth += (text[end] == "{") - (text[end] == "}")
        if depth == 0:
            return text[start:end]
    raise AssertionError(name)


def pairs(text, kind):
    pattern = (
        r"add_to_array\s*=\s*\{\s*array\s*=\s*global\.triggerable_scenario_"
        + kind + r"_ids\s+value\s*=\s*constant:triggerable_scenario_id\.([\w]+)\s*\}"
        r"\s*add_to_array\s*=\s*\{\s*array\s*=\s*global\.triggerable_scenario_"
        + kind + r"_name_sort_values\s+value\s*=\s*constant:triggerable_scenario_sort_value\.([\w]+)\s*\}"
    )
    result = re.findall(pattern, text)
    expected = dict(re.findall(pattern, (BASE / "effects.before.txt").read_text(encoding="utf-8-sig")))
    expected["assassin_network"] = "assassin_network_name"
    assert all(expected.get(name) == sort for name, sort in result)
    return [name for name, _ in result]


constants = read("common/script_constants/chaosx_triggerable_scenarios_constants.txt")
identities = dict(re.findall(r"(\w+)\s*=\s*(\d+)", block(constants, "triggerable_scenario_id")))
assert identities["assassin_network"] == "19"
assert identities["random_civil_war"] == "18"
assert "17" not in identities.values()
assert len(identities.values()) == len(set(identities.values()))
fallout = read("common/script_constants/fallout_consolidated_constants.txt")
assert re.search(r"triggerable_scenario_id\s*=\s*14", block(fallout, "fallout_manual_scenario_identity"))

effects = read("common/scripted_effects/chaosx_triggerable_scenarios_effects.txt")
before = (BASE / "effects.before.txt").read_text(encoding="utf-8-sig")
all_ids = pairs(block(effects, "triggerable_scenarios_initialize_registry"), "all")
old_ids = pairs(block(before, "triggerable_scenarios_initialize_registry"), "all")
assert len(all_ids) == 17 and len(set(all_ids)) == 17
assert [x for x in all_ids if x != "assassin_network"] == old_ids
views = pairs(block(effects, "triggerable_scenarios_rebuild_view"), "view")
old_views = pairs(block(before, "triggerable_scenarios_rebuild_view"), "view")
assert len(views) == 4 * len(all_ids)
for i in range(4):
    rows = views[i * 17:(i + 1) * 17]
    assert set(rows) == set(all_ids) and len(rows) == len(set(rows))
    assert [x for x in rows if x != "assassin_network"] == old_views[i * 16:(i + 1) * 16]
assert views[0] == "assassin_network" and views[33] == "assassin_network"
assert views[34] == "assassin_network" and views[-1] == "assassin_network"

localisation = read("localisation/english/039_murder_mystery_l_english.yml")
keys = set(re.findall(r"(?m)^([\w.]+):", localisation))
getters = read("common/scripted_localisation/chaosx_scripted_localisation_scenarios.txt")
required = ["GetTriggerableScenarioSelectedName", "GetTriggerableScenarioEntryName",
            "GetTriggerableScenarioEntryId", "GetTriggerableScenarioSelectedDesc",
            "GetTriggerableScenarioTypeName", "GetTriggerableScenarioWarning",
            "GetTriggerableScenarioLaunchStatus"]
for getter in required:
    start = getters.index("name = " + getter)
    next_start = getters.find("\ndefined_text =", start)
    section = getters[start:next_start if next_start != -1 else len(getters)]
    assert "triggerable_scenario_id.assassin_network" in section, getter
    for key in re.findall(r"localization_key\s*=\s*(chaosx\.scenarios\.[\w.]*assassin_network[\w.]*)", section):
        assert key in keys, key
for level in ["low", "medium", "high", "maximum"]:
    assert "chaosx.scenarios.assassin_network.impact." + level in getters
assert 'chaosx.scenarios.entry.id.assassin_network: "#019"' in localisation
assert getters.index("localization_key = chaosx.scenarios.launch_status.assassin_network.resolved") < getters.index("localization_key = chaosx.scenarios.launch_status.assassin_network.already_launched")

shared_gate = read("common/scripted_triggers/chaosx_triggerable_scenarios_triggers.txt")
assert "murder_mystery_triggerable_scenario_can_launch = yes" in block(shared_gate, "triggerable_scenario_can_launch_selected")
manual_gate = block(read("common/scripted_triggers/039_murder_mystery_integration_triggers.txt"), "murder_mystery_triggerable_scenario_can_launch")
assert "world_end_disabled" not in manual_gate and "murder_mystery_event_is_available" not in manual_gate
for flag in ["world_end", "murder_mystery_runtime_active", "murder_mystery_runtime_resolved", "murder_mystery_scenario_launch_in_progress", "murder_mystery_scenario_launched"]:
    assert "NOT = { has_global_flag = " + flag + " }" in manual_gate
assert "murder_mystery_has_any_scenario_host = yes" in manual_gate
terminal_gate = block(read("common/scripted_triggers/039_murder_mystery_integration_triggers.txt"), "murder_mystery_world_branch_is_enabled")
assert "NOT = { has_global_flag = world_end_disabled }" in terminal_gate
assert "world_end_murder_mystery_world_of_anarchy_scenario_enabled = yes" in terminal_gate
dispatch = block(effects, "trigger_selected_chaosx_scenario")
assert "triggerable_scenario_id.assassin_network" in dispatch
assert "murder_mystery_triggerable_scenario_launch = yes" in dispatch
gui = read("common/scripted_guis/chaosx_scripted_gui_settings.txt")
assert "trigger_selected_chaosx_scenario = yes" in gui
assert "global.triggerable_scenario_view_ids^triggerable_scenario_entry_index" in gui
assert gui.count("triggerable_scenario_can_launch_selected = yes") >= 2

integration = read("common/scripted_effects/039_murder_mystery_integration_effects.txt")
launch = block(integration, "murder_mystery_triggerable_scenario_launch")
assert "limit = { murder_mystery_triggerable_scenario_can_launch = yes }" in launch
assert "global.murder_mystery_scenario_launch_intensity = triggerable_scenarios_intensity" in launch
assert "murder_mystery_run_opening = yes" in launch
setup = block(integration, "murder_mystery_apply_scenario_setup")
assert re.search(r"else\s*=\s*\{\s*set_variable\s*=\s*\{\s*murder_mystery_scenario_stage_input\s*=\s*constant:murder_mystery_evolution.stage_i", setup)
root = read("events/039_murder_mystery.txt").split("country_event =", 2)[1]
assert "murder_mystery_run_opening = yes" in root
for cleanup in ["clr_global_flag = murder_mystery_scenario_launch_in_progress",
                "clear_variable = global.murder_mystery_scenario_launch_intensity"]:
    assert cleanup in launch
    assert cleanup in block(integration, "murder_mystery_run_opening")

workbook = load_workbook(ROOT / "docs/spreadsheets/chaos_redux_events_catalog.xlsx")
sheet = workbook["Scenarios"]
assert sheet["B14"].value == "Assassin Network" and sheet["A14"].value == "SCN-019"
assert sheet["C14"].value == re.search(r'(?m)^chaosx\.scenarios\.assassin_network\.desc: "(.*)"$', localisation)[1]
assert sheet["D14"].value == re.search(r'(?m)^chaosx\.scenarios\.type\.assassin_network: "(.*)"$', localisation)[1]
expected_warnings = "\n\n".join(re.search(
    r'(?m)^chaosx\.scenarios\.assassin_network\.impact\.' + level + r': "(.*)"$', localisation)[1]
    for level in ["low", "medium", "high", "maximum"])
assert sheet["E14"].value == expected_warnings
assert sheet["E14"].value.startswith("§R")
with (ROOT / "docs/spreadsheets/chaos_redux_scenarios_catalog.csv").open(encoding="utf-8-sig", newline="") as export:
    exported = next(row for row in csv.reader(export) if row[1] == "Assassin Network")
assert exported[0] == sheet["A14"].value and exported[3] == sheet["D14"].value
assert exported[4] == sheet["E14"].value

cluster_constants = read("common/script_constants/event_cluster_constants.txt")
cluster_ids = dict(re.findall(r"(\w+)\s*=\s*(\d+)", block(cluster_constants, "event_cluster_id")))
assert cluster_ids["intelligence"] == "20"
cluster_row = next(row for row in workbook["Clusters"].iter_rows(values_only=True) if row[1] == "Intelligence")
assert cluster_row[0] == 20
for event_id in [39, 52]:
    event_row = next(row for row in workbook["Events"].iter_rows(values_only=True) if row[0] == event_id)
    assert str(event_row[11]) == "20"
assert workbook["Cluster Memberships"]["A42"].value == 20
assert workbook["Cluster Memberships"]["A43"].value == 20
for name in ["Events", "Clusters", "Scenarios"]:
    ids = [row[0] for row in workbook[name].iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(ids) == len(set(ids)), name

print("PASS: 17 unique shared rows, four aligned 17-row views, preserved prior order/IDs, seven getters, four warnings, gate/dispatch/root linkage, Low fallback and launch-cache cleanup.")
print("PASS: workbook ID, single type, four intensity descriptions and regenerated scenario export match their runtime presentation.")
print("PASS: Intelligence ID 20, both member cross-links and workbook identity uniqueness agree.")
print("LIMIT: static source contracts only; country transactions, render, live launch and persistence remain separate evidence.")
